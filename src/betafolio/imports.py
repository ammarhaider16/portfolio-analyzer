"""Portfolio file imports use quantities; the system supplies market prices."""

import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Literal

from pydantic import Field, field_validator

from .models import Contract, PortfolioSnapshot, Position


class DraftPosition(Contract):
    symbol: str = Field(min_length=1, max_length=40)
    quantity: Decimal | None = Field(default=None, ge=0)
    price: Decimal | None = Field(default=None, ge=0)
    asset_type: str = Field(default="unknown", max_length=40)
    instrument_name: str | None = Field(default=None, max_length=200)
    exchange: str | None = Field(default=None, max_length=40)
    sector: str | None = Field(default=None, max_length=100)
    previous_close: Decimal | None = Field(default=None, ge=0)
    price_as_of: datetime | None = None
    price_session: date | None = None
    price_source: str | None = None

    @field_validator("symbol")
    @classmethod
    def valid_symbol(cls, value):
        return Position.normalize_symbol(value)


class CsvImportDraft(Contract):
    """An unsaved import awaits system prices and is never a portfolio snapshot."""

    as_of: datetime
    positions: list[DraftPosition] = Field(default_factory=list, max_length=1000)
    cash: Decimal = Decimal(0)
    source: Literal["csv"] = "csv"
    account_id: str = "manual"
    unsupported: list[str] = Field(default_factory=list)


def parse_amount(value: str) -> Decimal:
    """Accept plain decimals or conventional dollar amounts without losing precision."""
    value = value.strip()
    if value.startswith("$"):
        value = value[1:].strip()
    if not re.fullmatch(r"(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.\d+)?|\.\d+", value):
        raise ValueError("Amount must be a finite nonnegative number")
    return Decimal(value.replace(",", ""))


def validate_portfolio_snapshot(snapshot: PortfolioSnapshot) -> PortfolioSnapshot:
    """Normalize cash aliases and reject ambiguous duplicate holdings."""
    positions, seen, cash = [], set(), snapshot.cash
    for position in snapshot.positions:
        kind = position.asset_type.strip().lower()
        if kind == "cash" or (
            position.symbol in {"CASH", "USD"} and kind not in {"etf", "fund", "mutual_fund"}
        ):
            cash += position.quantity * (position.price if position.price is not None else Decimal(1))
            continue
        if position.symbol in seen:
            raise ValueError(f"Duplicate symbol {position.symbol}; combine lots first")
        seen.add(position.symbol)
        positions.append(
            position.model_copy(update={"asset_type": position.asset_type.strip().lower() or "unknown"})
        )
    if not positions and cash == 0:
        raise ValueError("Portfolio contains no holdings or cash")
    return snapshot.model_copy(
        update={
            "positions": positions,
            "cash": cash,
            "unsupported": [p.symbol for p in positions if p.asset_type not in {"stock", "etf", "unknown"}],
        }
    )


def parse_holdings_csv(content: str, as_of: datetime) -> CsvImportDraft:
    if len(content) > 5 * 1024 * 1024:
        raise ValueError("CSV limit is 5 MB")
    if as_of.tzinfo is None:
        raise ValueError("Snapshot date must include a timezone")
    reader = csv.DictReader(io.StringIO(content.lstrip("\ufeff")))
    aliases = {
        "ticker": "symbol",
        "shares": "quantity",
        "share count": "quantity",
        "current price": "price",
        "share price": "price",
    }
    reader.fieldnames = [
        aliases.get(name.strip().lower(), name.strip().lower()) for name in (reader.fieldnames or [])
    ]
    if not {"symbol", "quantity"}.issubset(reader.fieldnames):
        raise ValueError("Include ticker and quantity columns (symbol and shares also work)")
    if len(reader.fieldnames) != len(set(reader.fieldnames)):
        raise ValueError("CSV has duplicate columns after matching column names")
    positions, unsupported, cash, seen = [], [], Decimal(0), set()

    def number(value, field, line):
        try:
            result = parse_amount(value)
            if not result.is_finite() or result < 0:
                raise ValueError()
            return result
        except (InvalidOperation, ValueError, AttributeError):
            raise ValueError(f"Row {line}: {field} must be a finite nonnegative number") from None

    for line, row in enumerate(reader, 2):
        if None in row:
            raise ValueError(f"Row {line}: too many columns; quote values containing commas")
        if line > 1001:
            raise ValueError("CSV limit is 1,000 rows")
        symbol = (row.get("symbol") or "").strip().upper()
        if not symbol or len(symbol) > 32:
            raise ValueError(f"Row {line}: provide a symbol of at most 32 characters")
        quantity = number(row.get("quantity"), "quantity", line)
        kind = (row.get("asset_type") or "unknown").strip().lower()
        if kind == "cash" or (symbol in {"CASH", "USD"} and kind not in {"etf", "fund", "mutual_fund"}):
            cash += quantity
            continue
        if symbol in seen:
            raise ValueError(f"Row {line}: duplicate symbol {symbol}; combine lots first")
        seen.add(symbol)
        if kind not in {"stock", "etf", "unknown"}:
            unsupported.append(symbol)
        positions.append(
            DraftPosition(
                symbol=symbol,
                quantity=quantity,
                asset_type=kind,
                sector=(row.get("sector") or "").strip() or None,
            )
        )
    if not positions and cash == 0:
        raise ValueError("CSV contains no holdings or cash")
    return CsvImportDraft(as_of=as_of, positions=positions, cash=cash, unsupported=unsupported)


def parse_holdings_file(content: bytes, filename: str, as_of: datetime) -> CsvImportDraft:
    """Read a CSV or the first visible Excel worksheet as an editable holdings draft."""
    from pathlib import Path
    from zipfile import ZipFile

    if not content or len(content) > 5 * 1024 * 1024:
        raise ValueError("Choose a nonempty file up to 5 MB")
    extension = Path(filename).suffix.lower()
    if extension == ".csv":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("cp1252")
        return parse_holdings_csv(text, as_of)
    if extension not in {".xlsx", ".xls"}:
        raise ValueError("Choose a CSV, XLSX or XLS file")
    rows = []
    try:
        if extension == ".xlsx":
            # Reject oversized expanded archives before XML parsing. Formulas are
            # never evaluated; cached values are used, with blanks left for review.
            with ZipFile(io.BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries) > 1000 or sum(entry.file_size for entry in entries) > 20 * 1024 * 1024:
                    raise ValueError("This workbook is too large. Export its holdings sheet as CSV.")
            import openpyxl

            workbook = openpyxl.load_workbook(
                io.BytesIO(content), read_only=True, data_only=True, keep_links=False
            )
            try:
                sheet = next((sheet for sheet in workbook.worksheets if sheet.sheet_state == "visible"), None)
                if sheet is None:
                    raise ValueError("The workbook has no visible worksheet")
                if (sheet.max_column or 0) > 50:
                    raise ValueError("The holdings worksheet must have no more than 50 columns")
                rows = list(
                    sheet.iter_rows(
                        max_row=min(sheet.max_row or 1002, 1002),
                        max_col=min(sheet.max_column or 50, 50),
                        values_only=True,
                    )
                )
            finally:
                workbook.close()
        else:
            import xlrd

            workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
            try:
                sheet = next(
                    (
                        workbook.sheet_by_index(index)
                        for index in range(workbook.nsheets)
                        if not workbook.sheet_by_index(index).visibility
                    ),
                    None,
                )
                if sheet is None:
                    raise ValueError("The workbook has no visible worksheet")
                if sheet.ncols > 50:
                    raise ValueError("The holdings worksheet must have no more than 50 columns")
                rows = [sheet.row_values(index) for index in range(min(sheet.nrows, 1002))]
            finally:
                workbook.release_resources()
    except ValueError:
        raise
    except Exception:
        raise ValueError(
            "The workbook could not be read. Choose a valid Excel file or export it as CSV."
        ) from None
    # Ignore trailing formatted cells/rows, but keep missing values in real records.
    while rows and not any(value is not None and str(value).strip() for value in rows[-1]):
        rows.pop()
    if not rows:
        raise ValueError("The worksheet is empty")
    columns = len(rows[0])
    while columns and (rows[0][columns - 1] is None or not str(rows[0][columns - 1]).strip()):
        columns -= 1
    if not columns:
        raise ValueError("The first row needs ticker and quantity column names")
    if len(rows) > 1001:
        raise ValueError("The holdings worksheet is limited to 1,000 rows")
    stream = io.StringIO()
    writer = csv.writer(stream)
    for row in rows:
        writer.writerow(["" if value is None else str(value) for value in row[:columns]])
    return parse_holdings_csv(stream.getvalue(), as_of)
